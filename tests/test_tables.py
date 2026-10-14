import pytest
from openpyxl import Workbook

from optopus_studio.tables import TableError, preview, read_records


def write(tmp_path, name, text):
    path = tmp_path / name
    path.write_text(text)
    return path


@pytest.mark.parametrize(
    ("name", "text"),
    [
        ("stores.csv", "name,demand,x\nA,3,1.5\nB,,2\n"),
        ("stores.tsv", "name\tdemand\tx\nA\t3\t1.5\nB\t\t2\n"),
        ("stores.txt", "name;demand;x\nA;3;1.5\nB;;2\n"),
        ("stores.csv", "﻿name,demand,x\nA,3,1.5\n,,\nB,,2\n"),
    ],
)
def test_delimited_files_become_records_with_numbers(tmp_path, name, text):
    assert read_records(write(tmp_path, name, text)) == [
        {"name": "A", "demand": 3, "x": 1.5},
        {"name": "B", "demand": None, "x": 2},
    ]


def test_an_excel_sheet_is_read_by_name(tmp_path):
    workbook = Workbook()
    workbook.active.title = "notes"
    sheet = workbook.create_sheet("staff")
    for row in [("name", "hours"), ("Aoki", 8), ("Baba", 6.5)]:
        sheet.append(row)
    path = tmp_path / "staff.xlsx"
    workbook.save(path)
    assert read_records(path, "staff") == [
        {"name": "Aoki", "hours": 8},
        {"name": "Baba", "hours": 6.5},
    ]
    assert "sheets: ['notes', 'staff']" in preview(path, "staff")
    with pytest.raises(TableError, match="no sheet"):
        read_records(path, "missing")


def test_a_preview_shows_the_shape_and_the_first_rows(tmp_path):
    rows = "\n".join(f"{i},{i * i}" for i in range(25))
    text = preview(write(tmp_path, "squares.csv", f"n,square\n{rows}\n"))
    assert "25 rows, columns: ['n', 'square']" in text
    assert '{"n": 9, "square": 81}' in text
    assert "15 more rows" in text


@pytest.mark.parametrize(("name", "text"), [("dup.csv", "a,a\n1,2\n"), ("notes.md", "# hi\n")])
def test_unreadable_tables_are_rejected(tmp_path, name, text):
    with pytest.raises(TableError):
        read_records(write(tmp_path, name, text))
