"""Reading the files a user uploads (CSV, TSV, Excel, JSON) as tables of records.

The agent previews an upload, then imports it into instance.json under a key, so the data the
model reads is frozen into every revision with it and never has to pass through the agent's
context.
"""

from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from typing import Any

UPLOAD_SUFFIXES = (".csv", ".tsv", ".txt", ".xlsx", ".json")
PREVIEW_ROWS = 10


class TableError(ValueError):
    pass


def cell(value: Any) -> Any:
    """A number when the text is one, the text otherwise; empty cells become None."""
    if isinstance(value, str):
        text = value.strip()
        if text == "":
            return None
        for parse in (int, float):
            try:
                return parse(text)
            except ValueError:
                pass
        return text
    return value


def _records(header: list[Any], rows: list[list[Any]]) -> list[dict[str, Any]]:
    columns = [
        str(h).strip() if h not in (None, "") else f"column_{i + 1}" for i, h in enumerate(header)
    ]
    if len(set(columns)) != len(columns):
        raise TableError(f"the header repeats a column name: {columns}")
    records = []
    for row in rows:
        values = [cell(v) for v in row]
        if all(v is None for v in values):
            continue
        values += [None] * (len(columns) - len(values))
        records.append(dict(zip(columns, values[: len(columns)], strict=True)))
    return records


def _delimited(text: str, delimiter: str | None) -> list[dict[str, Any]]:
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",\t;").delimiter
        except csv.Error:
            delimiter = ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    if not rows:
        return []
    return _records(rows[0], rows[1:])


def sheet_names(path: Path) -> list[str]:
    if path.suffix.lower() != ".xlsx":
        return []
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def read_records(path: Path, sheet: str | None = None) -> Any:
    """The file as a list of records (one dict per row, keyed by the header), or the JSON value."""
    suffix = path.suffix.lower()
    if suffix not in UPLOAD_SUFFIXES:
        raise TableError(f"cannot read {path.name}; expected one of {UPLOAD_SUFFIXES}")
    if suffix == ".json":
        return json.loads(path.read_text())
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet is not None and sheet not in workbook.sheetnames:
                raise TableError(f"no sheet {sheet!r}; the sheets are {workbook.sheetnames}")
            worksheet = workbook[sheet or workbook.sheetnames[0]]
            rows = [list(r) for r in worksheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
        return _records(rows[0], rows[1:]) if rows else []
    text = path.read_bytes().decode("utf-8-sig")
    return _delimited(text, "\t" if suffix == ".tsv" else None)


def preview(path: Path, sheet: str | None = None) -> str:
    """What the agent reads before importing: shape, columns, and the first rows."""
    data = read_records(path, sheet)
    lines = [f"{path.name} ({path.stat().st_size} bytes)"]
    sheets = sheet_names(path)
    if sheets:
        lines.append(f"sheets: {sheets}; showing {sheet or sheets[0]!r}")
    if isinstance(data, list) and all(isinstance(r, dict) for r in data):
        columns = list(data[0]) if data else []
        lines.append(f"{len(data)} rows, columns: {columns}")
        lines += [json.dumps(r, ensure_ascii=False) for r in data[:PREVIEW_ROWS]]
        if len(data) > PREVIEW_ROWS:
            lines.append(f"... {len(data) - PREVIEW_ROWS} more rows")
    else:
        text = json.dumps(data, ensure_ascii=False)
        lines.append(text[:3000] + (f"... ({len(text)} characters)" if len(text) > 3000 else ""))
    return "\n".join(lines)
